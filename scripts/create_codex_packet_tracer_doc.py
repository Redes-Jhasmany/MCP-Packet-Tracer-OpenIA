from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = REPO_ROOT / "docs" / "codex-mcp-packet-tracer-pasos.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3F6FA")
    set_cell_border(cell, "D9D9D9", "6")
    set_cell_margins(cell, top=120, start=140, bottom=120, end=140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = "DejaVu Sans Mono"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "DejaVu Sans Mono")
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(31, 41, 55)
    doc.add_paragraph()


def add_step(doc, number, title, body, commands=None):
    p = doc.add_paragraph()
    p.style = "Heading 2"
    p.add_run(f"{number}  {title}")
    for paragraph in body:
        para = doc.add_paragraph(paragraph)
        para.style = "Body Text"
    if commands:
        add_code_block(doc, commands)


def style_document(doc):
    styles = doc.styles
    for style_name in ["Normal", "Body Text"]:
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
        style.font.size = Pt(10.8)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.08
        style.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    title_ppr = title._element.find(qn("w:pPr"))
    if title_ppr is not None:
        title_borders = title_ppr.find(qn("w:pBdr"))
        if title_borders is not None:
            title_ppr.remove(title_borders)

    for style_name, size in [("Heading 1", 15), ("Heading 2", 12.5)]:
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)


def add_summary_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    widths = [Inches(1.45), Inches(2.35), Inches(2.65)]
    headers = ["Elemento", "Ruta o comando", "Resultado esperado"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = text
        set_cell_shading(cell, "1F4E79")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
    rows = [
        ("Servidor MCP", "/home/jhasmany/Proyectos/CCNA/MCP-Packet-Tracer", "Contiene el MCP completo con herramientas pt_*"),
        ("Python del MCP", ".venv/bin/python -m packet_tracer_mcp --stdio", "Arranca el servidor para Codex"),
        ("Extensión PT", "releases/V5.2.pts", "Habilita MCP BUILDER dentro de Packet Tracer"),
        ("Puerto bridge", "http://127.0.0.1:54321", "Debe aparecer CONNECTED por HTTP"),
        ("Prueba", "pt_full_build", "Crea y verifica una topología real"),
    ]
    for idx, row_data in enumerate(rows, start=1):
        row = table.add_row()
        for i, text in enumerate(row_data):
            cell = row.cells[i]
            cell.text = text
            set_cell_border(cell)
            set_cell_margins(cell)
            if idx % 2 == 0:
                set_cell_shading(cell, "F6F8FB")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(9.2)
                    run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph()


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Pasos para conectar Codex con Packet Tracer MCP")
    title_ppr = title._p.get_or_add_pPr()
    title_borders = title_ppr.find(qn("w:pBdr"))
    if title_borders is not None:
        title_ppr.remove(title_borders)

    intro = doc.add_paragraph()
    intro.style = "Body Text"
    intro.add_run("Objetivo. ").bold = True
    intro.add_run(
        "Esta guía documenta el procedimiento usado para conectar Codex con Cisco Packet Tracer "
        "mediante el servidor MCP completo de Packet Tracer. Al terminar, Codex puede usar "
        "herramientas pt_* como pt_bridge_status, pt_list_devices, pt_full_build y pt_query_topology "
        "para crear, desplegar y leer topologías reales dentro de Packet Tracer."
    )

    p = doc.add_paragraph()
    p.style = "Body Text"
    p.add_run("Resultado comprobado. ").bold = True
    p.add_run(
        "La prueba final desplegó una topología con un router, un switch y dos PCs. "
        "El bridge reportó conexión HTTP en 127.0.0.1:54321, verificó 4 de 4 dispositivos "
        "y 3 de 3 enlaces, y luego pt_query_topology leyó la topología viva desde Packet Tracer."
    )

    doc.add_paragraph("Resumen de la configuración", style="Heading 1")
    add_summary_table(doc)

    doc.add_paragraph("Procedimiento", style="Heading 1")
    add_step(
        doc,
        1,
        "Preparar el proyecto local",
        [
            "Clona el repositorio oficial del servidor MCP completo en una carpeta estable. "
            "En este equipo se usó la ruta de proyectos de CCNA para separar el MCP completo del script mínimo anterior."
        ],
        [
            "git clone https://github.com/Mats2208/MCP-Packet-Tracer.git /home/jhasmany/Proyectos/CCNA/MCP-Packet-Tracer",
            "cd /home/jhasmany/Proyectos/CCNA/MCP-Packet-Tracer",
        ],
    )

    add_step(
        doc,
        2,
        "Crear el entorno Python e instalar el MCP",
        [
            "Usa un entorno virtual local para no modificar Python global. "
            "El paquete instalado expone el módulo packet_tracer_mcp y las herramientas completas pt_*."
        ],
        [
            "python3 -m venv /home/jhasmany/Proyectos/CCNA/MCP-Packet-Tracer/.venv",
            "/home/jhasmany/Proyectos/CCNA/MCP-Packet-Tracer/.venv/bin/python -m pip install --upgrade pip",
            "/home/jhasmany/Proyectos/CCNA/MCP-Packet-Tracer/.venv/bin/pip install -e /home/jhasmany/Proyectos/CCNA/MCP-Packet-Tracer",
        ],
    )

    add_step(
        doc,
        3,
        "Registrar el servidor en Codex",
        [
            "Este es el equivalente en Codex al comando claude mcp add de la guía de Claude. "
            "Primero elimina cualquier servidor packet-tracer anterior si apuntaba al script mínimo, y luego registra el MCP completo."
        ],
        [
            "codex mcp remove packet-tracer",
            "codex mcp add packet-tracer -- /home/jhasmany/Proyectos/CCNA/MCP-Packet-Tracer/.venv/bin/python -m packet_tracer_mcp --stdio",
            "codex mcp get packet-tracer",
        ],
    )

    add_step(
        doc,
        4,
        "Instalar la extensión en Packet Tracer",
        [
            "Descarga la extensión V5.2.pts del release correspondiente y agrégala desde la ventana de módulos de scripting de Packet Tracer. "
            "Después de agregarla, selecciona MCP BUILDER o MCP Control Center y pulsa Start."
        ],
        [
            "mkdir -p /home/jhasmany/Proyectos/CCNA/MCP-Packet-Tracer/releases",
            "curl -L --fail -o /home/jhasmany/Proyectos/CCNA/MCP-Packet-Tracer/releases/V5.2.pts https://github.com/Mats2208/MCP-Packet-Tracer/releases/download/v0.9.0/V5.2.pts",
        ],
    )

    add_step(
        doc,
        5,
        "Abrir el panel MCP BUILDER",
        [
            "En Cisco Packet Tracer abre Extensions > MCP BUILDER y deja la ventana abierta. "
            "El panel se conecta al bridge local cuando el servidor MCP está vivo. Si la ventana se cierra, el canal por archivo puede tomar el relevo, pero para la primera prueba es mejor mantenerla abierta."
        ],
    )

    add_step(
        doc,
        6,
        "Verificar la conexión",
        [
            "En una tarea nueva de Codex, las herramientas pt_* deben estar disponibles. "
            "La verificación correcta de pt_bridge_status debe mostrar CONNECTED por HTTP con la URL 127.0.0.1:54321."
        ],
        [
            "pt_bridge_status",
            "Resultado esperado: CONNECTED por HTTP (ventana MCP Control Center abierta) -- http://127.0.0.1:54321",
        ],
    )

    add_step(
        doc,
        7,
        "Ejecutar una prueba real",
        [
            "Antes de crear una topología, llama a pt_list_devices para confirmar modelos y puertos disponibles. "
            "Luego usa pt_full_build para desplegar una red pequeña y pt_query_topology para leer el resultado desde Packet Tracer."
        ],
        [
            "pt_list_devices",
            "pt_full_build routers=1 pcs_per_lan=2 switches_per_router=1 dhcp=true routing=none router_model=2911 switch_model=2960-24TT template=single_lan deploy=true",
            "pt_query_topology",
        ],
    )

    doc.add_paragraph("Resultado de la prueba realizada", style="Heading 1")
    checks = [
        "El bridge conectó por HTTP en 127.0.0.1:54321.",
        "pt_full_build creó R1, SW1, PC1 y PC2.",
        "La validación marcó PASS.",
        "El despliegue envió 11 comandos a Packet Tracer.",
        "Packet Tracer verificó 4 de 4 dispositivos y 3 de 3 enlaces.",
        "pt_query_topology confirmó que la topología viva contenía R1, SW1, PC1 y PC2.",
    ]
    for item in checks:
        p = doc.add_paragraph(style="Body Text")
        p.style = "List Bullet"
        p.add_run(item)

    doc.add_paragraph("Notas de solución de problemas", style="Heading 1")
    trouble = [
        ("Codex solo muestra tres herramientas packet_tracer_*", "La tarea actual cargó el MCP viejo. Abre una tarea nueva de Codex o reinicia la sesión para cargar las herramientas pt_*."),
        ("pt_bridge_status dice que no hay canal", "Abre Extensions > MCP BUILDER y espera unos segundos. El servidor MCP debe estar vivo para que el panel haga polling."),
        ("El puerto 54321 no aparece", "El bridge HTTP se enciende cuando se invoca una herramienta como pt_bridge_status desde el MCP completo."),
        ("Se despliega en modo manual", "Significa que el bridge no estaba conectado al momento de ejecutar pt_full_build. Repite la prueba después de ver CONNECTED por HTTP."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for i, h in enumerate(["Situación", "Acción recomendada"]):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, "404040")
        set_cell_border(cell)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)
    for idx, (a, b) in enumerate(trouble, start=1):
        row = table.add_row()
        for i, text in enumerate([a, b]):
            cell = row.cells[i]
            cell.text = text
            set_cell_border(cell)
            set_cell_margins(cell)
            if idx % 2 == 0:
                set_cell_shading(cell, "F7F7F7")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9.2)

    doc.add_section(WD_SECTION.CONTINUOUS)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))


if __name__ == "__main__":
    main()
